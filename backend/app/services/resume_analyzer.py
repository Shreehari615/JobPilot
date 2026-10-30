"""
Resume NLP Analyzer – Pipeline:
  Text Extraction → Cleaning → Section Detection →
  NER (spaCy) → Skill/Edu/Exp extraction (regex+taxonomy) →
  Candidate Representation

Works entirely without an LLM.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.ai.skill_taxonomy import extract_skills_from_text, normalize_skill
from app.utils.logger import get_logger

logger = get_logger(__name__)

# ── Lazy-load heavy NLP models ────────────────────────────────────────
_nlp = None


def _get_nlp():
    global _nlp
    if _nlp is None:
        try:
            import spacy
            _nlp = spacy.load("en_core_web_sm")
            logger.info("spaCy model loaded: en_core_web_sm")
        except Exception as e:
            logger.warning(f"spaCy load failed ({e}), using regex-only mode")
            _nlp = False
    return _nlp if _nlp else None


# ── Section header patterns ───────────────────────────────────────────
SECTION_PATTERNS = {
    "summary": re.compile(
        r"(?i)^(career\s+objective|professional\s+summary|career\s+summary|executive\s+summary|summary|objective|profile|about\s+me)\b",
        re.IGNORECASE,
    ),
    "skills": re.compile(
        r"(?i)(technical\s+)?skills?|competencies|technologies|tech\s+stack|tools?\s+&\s+technologies?",
        re.IGNORECASE,
    ),
    "education": re.compile(
        r"(?i)education|academic|qualification|degree|university|college|school",
        re.IGNORECASE,
    ),
    "experience": re.compile(
        r"(?i)^(work\s+experience|professional\s+experience|employment\s+history|work\s+history|professional\s+background|work\s+background|experience|employment|career\s+history)\b",
        re.IGNORECASE,
    ),
    "projects": re.compile(r"(?i)projects?|portfolio|work\s+samples?", re.IGNORECASE),
    "interests": re.compile(r"(?i)^(interests?|hobbies|extracurricular(\s+activities)?|activities)\b", re.IGNORECASE),
    "certifications": re.compile(
        r"(?i)certifications?|certificates?|licenses?|credentials?",
        re.IGNORECASE,
    ),
}


# ── Regex patterns ────────────────────────────────────────────────────
EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(
    r"(?:\+91[\s\-]?)?[789]\d{9}|(?:\+\d{1,3}[\s\-]?)?\(?\d{3}\)?[\s\-]\d{3}[\s\-]\d{4}"
)
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
DATE_RANGE_RE = re.compile(
    r"(?P<start>(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[\w]*\.?\s*\d{4}|"
    r"\d{4})\s*[-–—to]+\s*(?P<end>(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[\w]*\.?\s*\d{4}|"
    r"\d{4}|Present|Current|Now)",
    re.IGNORECASE,
)

DEGREE_PATTERNS = {
    "PhD": re.compile(r"\b(ph\.?d|doctor\s+of\s+philosophy)\b", re.IGNORECASE),
    "MBA": re.compile(r"\b(m\.?b\.?a|master\s+of\s+business)\b", re.IGNORECASE),
    "MCA": re.compile(r"\b(m\.?c\.?a|master\s+of\s+computer\s+application[s]?)\b", re.IGNORECASE),
    "MTech": re.compile(r"\b(m\.?\s*tech|master\s+of\s+technology)\b", re.IGNORECASE),
    "MSc": re.compile(r"\b(m\.?\s*sc|master\s+of\s+science)\b", re.IGNORECASE),
    "ME": re.compile(r"\b(m\.?\s*e\.?|master\s+of\s+engineering)\b", re.IGNORECASE),
    "BCA": re.compile(r"\b(b\.?c\.?a|bachelor\s+of\s+computer\s+application[s]?)\b", re.IGNORECASE),
    "BTech": re.compile(r"\b(b\.?\s*tech|bachelor\s+of\s+technology)\b", re.IGNORECASE),
    "BE": re.compile(r"\b(b\.?\s*e\.?|bachelor\s+of\s+engineering)\b", re.IGNORECASE),
    "BSc": re.compile(r"\b(b\.?\s*sc|bachelor\s+of\s+science)\b", re.IGNORECASE),
    "BA": re.compile(r"\b(b\.?\s*a\.?|bachelor\s+of\s+arts)\b", re.IGNORECASE),
    "BCom": re.compile(r"\b(b\.?\s*com|bachelor\s+of\s+commerce)\b", re.IGNORECASE),
    "BBA": re.compile(r"\b(b\.?b\.?a|bachelor\s+of\s+business)\b", re.IGNORECASE),
    "Diploma": re.compile(r"\b(diploma|polytechnic)\b", re.IGNORECASE),
    "PUC": re.compile(r"\b(puc|12th|class\s+xii|higher\s+secondary|intermediate|pre-university|plus\s+two)\b", re.IGNORECASE),
    "SSLC": re.compile(r"\b(sslc|10th|class\s+x|secondary|matriculation)\b", re.IGNORECASE),
}

DEGREE_LEVEL = {
    "PhD": 7, "MBA": 6, "MTech": 6, "ME": 6, "MSc": 6, "MCA": 6,
    "BTech": 5, "BE": 5, "BSc": 5, "BCA": 5, "BA": 5, "BCom": 5, "BBA": 5,
    "Diploma": 3, "PUC": 2, "SSLC": 1,
}


@dataclass
class CandidateProfile:
    """Structured representation of a candidate extracted from resume."""
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    skills: List[str] = field(default_factory=list)
    education: List[Dict[str, Any]] = field(default_factory=list)
    experience: List[Dict[str, Any]] = field(default_factory=list)
    certifications: List[str] = field(default_factory=list)
    projects: List[Any] = field(default_factory=list)
    job_titles: List[str] = field(default_factory=list)
    total_experience_years: float = 0.0
    highest_degree: Optional[str] = None
    highest_degree_level: int = 0
    summary: Optional[str] = None
    raw_text: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "email": self.email,
            "phone": self.phone,
            "skills": self.skills,
            "education": self.education,
            "experience": self.experience,
            "certifications": self.certifications,
            "projects": self.projects,
            "job_titles": self.job_titles,
            "total_experience_years": self.total_experience_years,
            "highest_degree": self.highest_degree,
            "highest_degree_level": self.highest_degree_level,
            "summary": self.summary,
        }


# ── Text Extraction ───────────────────────────────────────────────────
def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
        text_parts = []
        with fitz.open(file_path) as doc:
            for page in doc:
                text_parts.append(page.get_text())
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        raise ValueError(f"Could not extract text from PDF: {e}")


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise ValueError(f"Could not extract text from DOCX: {e}")


def extract_text(file_path: str, file_type: str) -> str:
    """Route to appropriate extractor."""
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


# ── Text Cleaning ─────────────────────────────────────────────────────
def clean_text(text: str) -> str:
    """Clean extracted resume text."""
    import unicodedata
    text = unicodedata.normalize("NFKD", text)
    # Convert range replacement chars (like 9095% or 2022  2025)
    text = re.sub(r"(\d+)\s*[\ufffd\u2013\u2014]\s*(\d+)", r"\1 - \2", text)
    # Normalize standalone bullet characters on their own lines
    text = re.sub(r"[\ufffd\u2022\u25cf\u25aa\u25b6\xb7\u2013\u2014]\s*\n\s*", "\n• ", text)
    text = text.replace("\ufffd", "-")
    # Remove null bytes and control chars
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)
    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse multiple blanks
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


# ── Section Detection ─────────────────────────────────────────────────
def detect_sections(text: str) -> Dict[str, str]:
    """
    Split resume text into named sections by detecting headers.
    Returns dict: section_name → section_text
    """
    lines = text.split("\n")
    sections: Dict[str, str] = {"header": ""}
    current_section = "header"
    current_lines: List[str] = []

    for line in lines:
        stripped = line.strip()
        matched_section = None
        # Check if this line is a section header
        for section_name, pattern in SECTION_PATTERNS.items():
            if pattern.match(stripped) and len(stripped) < 60:
                matched_section = section_name
                break

        if matched_section:
            sections[current_section] = "\n".join(current_lines)
            current_section = matched_section
            current_lines = []
        else:
            current_lines.append(line)

    sections[current_section] = "\n".join(current_lines)
    return sections


# ── Entity Extraction ─────────────────────────────────────────────────
def extract_email(text: str) -> Optional[str]:
    match = EMAIL_RE.search(text)
    return match.group() if match else None


def extract_phone(text: str) -> Optional[str]:
    match = PHONE_RE.search(text)
    return match.group() if match else None


def extract_name(text: str, header_text: str) -> Optional[str]:
    """
    Try spaCy PERSON entity first, fall back to first non-contact line.
    """
    nlp = _get_nlp()
    if nlp:
        doc = nlp(header_text[:500])
        for ent in doc.ents:
            if ent.label_ == "PERSON" and len(ent.text.split()) >= 2:
                return ent.text.strip()

    # Fallback: first non-empty, non-contact, non-URL line
    for line in header_text.split("\n")[:10]:
        line = line.strip()
        if (line and len(line) > 3 and len(line) < 60
                and not EMAIL_RE.search(line)
                and not PHONE_RE.search(line)
                and "http" not in line.lower()
                and "@" not in line):
            # Basic name validation: mostly alpha words
            words = line.split()
            if 1 <= len(words) <= 5 and all(w[0].isupper() or w.startswith("+") for w in words if w):
                return line
    return None


def extract_education(sections: Dict[str, str]) -> Tuple[List[Dict[str, Any]], Optional[str], int]:
    """
    Extract structured education entries line-by-line.
    Returns: (education_list, highest_degree, degree_level)
    """
    edu_text = sections.get("education", "").strip()
    if not edu_text:
        edu_text = sections.get("header", "")

    education: List[Dict[str, Any]] = []
    highest_degree = None
    highest_level = 0

    lines = [l.strip() for l in edu_text.split("\n") if l.strip()]

    degree_priority = [
        ("PhD", DEGREE_PATTERNS["PhD"]),
        ("MTech", DEGREE_PATTERNS["MTech"]),
        ("ME", DEGREE_PATTERNS["ME"]),
        ("MCA", DEGREE_PATTERNS["MCA"]),
        ("MBA", DEGREE_PATTERNS["MBA"]),
        ("MSc", DEGREE_PATTERNS["MSc"]),
        ("BTech", DEGREE_PATTERNS["BTech"]),
        ("BE", DEGREE_PATTERNS["BE"]),
        ("BCA", DEGREE_PATTERNS["BCA"]),
        ("BSc", DEGREE_PATTERNS["BSc"]),
        ("BCom", DEGREE_PATTERNS["BCom"]),
        ("BBA", DEGREE_PATTERNS["BBA"]),
        ("BA", DEGREE_PATTERNS["BA"]),
        ("Diploma", DEGREE_PATTERNS["Diploma"]),
        ("PUC", DEGREE_PATTERNS["PUC"]),
        ("SSLC", DEGREE_PATTERNS["SSLC"]),
    ]

    for line in lines:
        line_clean = line.replace("\ufffd", "-").strip()
        matched_deg = None
        for deg_name, deg_re in degree_priority:
            if deg_re.search(line_clean):
                matched_deg = deg_name
                break

        has_inst_kw = bool(re.search(r"(?i)\b(college|university|institute|school|academy|vidyalaya|campus)\b", line_clean))
        if not matched_deg and not has_inst_kw:
            # Check if this line is a continuation line containing graduation year/years for the previous degree
            if education and not education[-1].get("graduation_year"):
                cont_years = YEAR_RE.findall(line_clean)
                if cont_years:
                    education[-1]["graduation_year"] = int(cont_years[-1])
                    education[-1]["year"] = f"{cont_years[0]}-{cont_years[1]}" if len(cont_years) >= 2 else str(cont_years[0])
            continue

        # Extract years / year range
        years = YEAR_RE.findall(line_clean)
        grad_year = None
        year_str = None
        if len(years) >= 2:
            year_str = f"{years[0]}-{years[1]}"
            grad_year = int(years[1])
        elif len(years) == 1:
            year_str = str(years[0])
            grad_year = int(years[0])

        # Extract grade / CGPA / percentage
        grade = None
        cgpa_m = re.search(r"(?:CGPA|Percentage|Score|Grade)\s*[:=]?\s*([\d\.]+(?:\s*%)?)", line_clean, re.IGNORECASE)
        if cgpa_m:
            grade = cgpa_m.group(0).strip()

        # Extract institution by removing known degree, grade, years, and status
        inst = line_clean
        if matched_deg:
            inst = re.sub(r"^\s*" + re.escape(matched_deg) + r"[\s,–\-:]*", "", inst, flags=re.IGNORECASE)
        if grade:
            inst = inst.replace(grade, "")
        inst = re.sub(r"\b(?:19|20)\d{2}(?:\s*[-–—to]+\s*(?:19|20)\d{2})?\b", "", inst)
        inst = re.sub(r"\([^)]*\)", "", inst)
        inst = re.sub(r"\s*\|\s*", ", ", inst)
        inst = re.sub(r"\s*,\s*", ", ", inst)
        inst = inst.strip(" ,|–-")

        if matched_deg or (inst and len(inst) > 3) or grad_year:
            deg = matched_deg or ("Degree" if has_inst_kw else None)
            education.append({
                "degree": deg,
                "institution": inst if len(inst) > 3 else "University / Institution",
                "year": year_str or (str(grad_year) if grad_year else None),
                "graduation_year": grad_year,
                "grade": grade,
                "field_of_study": grade or None,
            })
            if matched_deg:
                level = DEGREE_LEVEL.get(matched_deg, 0)
                if level > highest_level:
                    highest_level = level
                    highest_degree = matched_deg

    # Fallback to global regex if line-by-line produced no results
    if not education:
        all_text = "\n".join(sections.values())
        for degree_name, pattern in degree_priority:
            if pattern.search(all_text):
                level = DEGREE_LEVEL.get(degree_name, 0)
                if level > highest_level:
                    highest_level = level
                    highest_degree = degree_name

        years = YEAR_RE.findall(edu_text or all_text)
        grad_year = max([int(y) for y in years], default=None) if years else None
        if highest_degree or grad_year:
            education.append({
                "degree": highest_degree or "Degree",
                "institution": "University / College",
                "year": str(grad_year) if grad_year else None,
                "graduation_year": grad_year,
                "grade": None,
                "field_of_study": None,
            })

    return education, highest_degree, highest_level


def extract_experience(sections: Dict[str, str]) -> Tuple[List[Dict], List[str], float]:
    """
    Extract work experience entries.
    Returns: (experience_list, job_titles, total_years)
    """
    exp_text = sections.get("experience", "")
    experiences = []
    job_titles = []
    total_months = 0

    # Common job title patterns
    title_patterns = [
        r"(?i)\b(software\s+engineer|developer|programmer|architect|analyst|data\s+scientist|"
        r"ml\s+engineer|backend\s+developer|frontend\s+developer|full\s+stack|"
        r"devops\s+engineer|cloud\s+engineer|sre|tech\s+lead|senior\s+engineer|"
        r"junior\s+engineer|intern|consultant|manager|project\s+manager|"
        r"product\s+manager|scrum\s+master|data\s+engineer|data\s+analyst|"
        r"research\s+engineer|research\s+scientist|ai\s+engineer)\b"
    ]

    for pattern in title_patterns:
        for match in re.finditer(pattern, exp_text):
            title = match.group().strip().title()
            if title not in job_titles:
                job_titles.append(title)

    # Extract date ranges for duration
    for match in DATE_RANGE_RE.finditer(exp_text):
        start_str = match.group("start")
        end_str = match.group("end")
        months = _estimate_months(start_str, end_str)
        total_months += months
        experiences.append({
            "start": start_str,
            "end": end_str,
            "duration_months": months,
            "job_title": job_titles[0] if job_titles else None,
            "company": None,
        })

    # If no date ranges found, look for year patterns
    if not experiences and exp_text:
        years_found = [int(y) for y in YEAR_RE.findall(exp_text)]
        if len(years_found) >= 2:
            span = max(years_found) - min(years_found)
            total_months = span * 12
            if span > 0:
                experiences.append({
                    "start": str(min(years_found)),
                    "end": str(max(years_found)),
                    "duration_months": span * 12,
                    "job_title": job_titles[0] if job_titles else None,
                    "company": None,
                })

    total_years = round(total_months / 12, 1)
    return experiences, job_titles, total_years


def _estimate_months(start: str, end: str) -> int:
    """Estimate months between two date strings."""
    import datetime

    def parse_year_month(s: str):
        s = s.strip()
        if s.lower() in ("present", "current", "now"):
            now = datetime.date.today()
            return now.year, now.month
        # Try year only
        if re.match(r"^\d{4}$", s):
            return int(s), 1
        # Try month year
        month_map = {
            "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
            "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
        }
        for abbr, num in month_map.items():
            if abbr in s.lower():
                year_match = YEAR_RE.search(s)
                year = int(year_match.group()) if year_match else datetime.date.today().year
                return year, num
        year_match = YEAR_RE.search(s)
        if year_match:
            return int(year_match.group()), 1
        return None, None

    sy, sm = parse_year_month(start)
    ey, em = parse_year_month(end)
    if sy and ey:
        return max(0, (ey - sy) * 12 + (em - sm))
    return 0


def extract_certifications(sections: Dict[str, str]) -> List[str]:
    """Extract certification names from the certifications section."""
    cert_text = sections.get("certifications", "")
    certs = []
    # Common cert patterns
    cert_patterns = [
        r"(?i)(aws\s+certified[\w\s]+)",
        r"(?i)(google\s+cloud[\w\s]+)",
        r"(?i)(azure\s+[\w\s]+)",
        r"(?i)(cissp|ceh|ccna|ccnp|rhce|pmp|scrum\s+master|csa|cka|ckad)",
        r"(?i)(tensorflow\s+developer|tensorflow\s+certificate)",
        r"(?i)(oracle\s+certified[\w\s]+)",
        r"(?i)(microsoft\s+certified[\w\s]+)",
    ]
    for pattern in cert_patterns:
        for match in re.finditer(pattern, cert_text):
            cert = match.group().strip()
            if cert and cert not in certs:
                certs.append(cert)
    # Also grab bullet-pointed lines
    for line in cert_text.split("\n"):
        line = line.strip().lstrip("•·-*").strip()
        if 5 < len(line) < 100 and line not in certs:
            certs.append(line)
    return certs[:10]


def extract_projects(sections: Dict[str, str]) -> List[Dict[str, Any]]:
    """Extract structured projects from projects section."""
    proj_text = sections.get("projects", "").strip()
    if not proj_text:
        return []

    # Cut off interests / hobbies if caught in projects section
    for marker in ["INTERESTS", "Interests", "HOBBIES", "Hobbies", "ACTIVITIES", "Activities",
                   "EXTRA-CURRICULAR", "Extra-Curricular", "ACHIEVEMENTS", "Achievements"]:
        if marker in proj_text:
            proj_text = proj_text.split(marker)[0]

    # Normalize standalone bullets on their own lines
    proj_text = re.sub(r"[\ufffd\u2022\u25cf\u25aa\u25b6\xb7\u2013\u2014]\s*\n\s*", "\n• ", proj_text)
    proj_text = proj_text.replace("\ufffd", "-")

    lines = [l.strip() for l in proj_text.split("\n") if l.strip()]
    projects: List[Dict[str, Any]] = []
    current_proj: Optional[Dict[str, Any]] = None

    def _is_project_title(line: str) -> bool:
        """Return True only if the line looks like a project heading, not a sentence."""
        # Explicit bullet markers → never a title
        if line.startswith(("•", "-", "*", "–", "—")):
            return False
        # URL lines → never a title
        if line.lower().startswith("http") or "www." in line.lower():
            return False
        # Ends with colon → strong title signal
        if line.endswith(":"):
            return len(line) < 120
        # Ends with period or comma → it's a sentence / description line, not a title
        if line.endswith((".","," ,";")):
            return False
        # Long lines are descriptions, not titles
        if len(line) > 80:
            return False
        # Lines that start with a lowercase letter are descriptions
        if line and line[0].islower():
            return False
        # Lines that look like plain sentences (verb phrases): start with common action verbs
        action_verbs = re.compile(
            r"^(Built|Developed|Implemented|Designed|Created|Integrated|Reduced|Improved|"
            r"Automated|Trained|Deployed|Optimized|Managed|Led|Collaborated|Used|Worked|"
            r"Achieved|Processed|Configured|Set\s+up|Migrated|Analyzed|Handled)\b",
            re.IGNORECASE,
        )
        if action_verbs.match(line):
            return False
        # Word count check: titles are typically 2–9 words
        word_count = len(line.split())
        if word_count < 2 or word_count > 9:
            return False
        # Must have at least one word that is Title-cased or ALL-CAPS
        has_title_word = any(w[0].isupper() for w in line.split() if len(w) > 1)
        if not has_title_word:
            return False
        return True

    for line in lines:
        if _is_project_title(line) and len(line) > 5:
            # Clean title (strip trailing colon, project link refs)
            title = re.sub(r":?\s*Project\s*[Ll]ink:?.*", "", line, flags=re.IGNORECASE)
            title = title.rstrip(":").strip()
            if not title:
                continue
            current_proj = {
                "title": title,
                "description": "",
                "highlights": [],
                "skills": [],
            }
            projects.append(current_proj)
        elif current_proj:
            bullet = re.sub(r"^[•\-\*\–\—\s]+", "", line).strip()
            if bullet:
                current_proj["highlights"].append(bullet)
                if not current_proj["description"]:
                    current_proj["description"] = bullet
                else:
                    current_proj["description"] += " " + bullet
        else:
            # No current project yet — treat first non-title line as an orphan title
            # only if it really looks like a project name (short, capitalized)
            stripped = re.sub(r"^[•\-\*\s]+", "", line).strip()
            if stripped and len(stripped) < 70 and stripped[0].isupper():
                current_proj = {
                    "title": stripped.rstrip(":"),
                    "description": "",
                    "highlights": [],
                    "skills": [],
                }
                projects.append(current_proj)

    # Remove projects that are likely mis-parsed (title reads like a sentence, without project nouns)
    bad_title_re = re.compile(
        r"^(Built|Developed|Implemented|Designed|Created|Integrated|Reduced|Improved|"
        r"Automated|Trained|Deployed|Optimized|Managed|Led|Collaborated|Used|Worked|"
        r"Achieved|Processed|Configured|Set\s+up|Migrated|Analyzed|Handled)\b",
        re.IGNORECASE,
    )
    project_noun_re = re.compile(
        r"\b(System|Platform|Application|App|Dashboard|Tracker|Detector|Detection|Recognition|Model|Pipeline|Architecture|Tool|Engine|Network|Interface|Project)\b",
        re.IGNORECASE,
    )
    projects = [
        p for p in projects
        if not bad_title_re.match(p.get("title", "")) or project_noun_re.search(p.get("title", ""))
    ]

    # Extract technologies for each project
    for p in projects:
        full_text = p["title"] + " " + " ".join(p["highlights"])
        p["skills"] = extract_skills_from_text(full_text)

    # Fallback if no titles matched
    if not projects and proj_text:
        for line in lines[:6]:
            c = line.lstrip("•-* ").strip()
            if len(c) > 10:
                projects.append({
                    "title": c[:60],
                    "description": c,
                    "highlights": [],
                    "skills": extract_skills_from_text(c),
                })

    return projects[:8]


# ── Main Analyzer ─────────────────────────────────────────────────────
def analyze_resume(file_path: str, file_type: str) -> CandidateProfile:
    """
    Full resume analysis pipeline:
      extract → clean → section detect → NER → skill/edu/exp extraction
    """
    logger.info(f"Analyzing resume: {file_path} ({file_type})")

    # 1. Extract raw text
    raw_text = extract_text(file_path, file_type)

    # 2. Clean text
    cleaned = clean_text(raw_text)

    # 3. Detect sections
    sections = detect_sections(cleaned)

    # 4. Extract contact info
    email = extract_email(cleaned)
    phone = extract_phone(cleaned)
    name = extract_name(cleaned, sections.get("header", cleaned[:500]))

    # 5. Extract skills from technical and experience sections (avoid education/contact/interests)
    relevant_sections = [
        sections.get("skills", ""),
        sections.get("projects", ""),
        sections.get("experience", ""),
        sections.get("summary", ""),
        sections.get("certifications", ""),
    ]
    skill_text = "\n".join(s for s in relevant_sections if s.strip())
    if not skill_text.strip():
        # Fallback if no sections were identified
        skill_text = cleaned
    raw_skills = extract_skills_from_text(skill_text)


    # 6. Extract education
    education, highest_degree, degree_level = extract_education(sections)

    # 7. Extract experience
    experience, job_titles, total_years = extract_experience(sections)

    # 8. Extract certifications and projects
    certifications = extract_certifications(sections)
    projects = extract_projects(sections)

    # 9. Extract summary
    summary_text = sections.get("summary", "").strip()
    summary = summary_text[:500] if summary_text else None

    profile = CandidateProfile(
        name=name,
        email=email,
        phone=phone,
        skills=raw_skills,
        education=education,
        experience=experience,
        certifications=certifications,
        projects=projects,
        job_titles=job_titles,
        total_experience_years=total_years,
        highest_degree=highest_degree,
        highest_degree_level=degree_level,
        summary=summary,
        raw_text=cleaned[:5000],  # Store truncated version
    )

    logger.info(
        f"Resume analyzed: {len(raw_skills)} skills, "
        f"{len(education)} edu entries, "
        f"{total_years} years exp, "
        f"degree={highest_degree}"
    )
    return profile
